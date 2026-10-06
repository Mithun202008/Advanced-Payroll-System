import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')
doc = docx.Document('pbl/RiskLens ( SYAMRIT , S HARSHAVARTHANAN ).docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")
print(f"Total sections: {len(doc.sections)}")

# Let's inspect paragraphs 0 to 200 with their XML elements, page breaks, sectPr, styles
page_num = 1
for i, p in enumerate(doc.paragraphs[:205]):
    # check if there is a break
    xml = p._p.xml
    has_page_break = ('w:br w:type="page"' in xml) or ('lastRenderedPageBreak' in xml)
    has_sectPr = ('w:sectPr' in xml)
    
    # check runs
    txt = p.text.strip()
    
    pPr = p._p.find('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pPr')
    jc = p._p.xpath('./w:pPr/w:jc/@w:val')
    sb = p._p.xpath('./w:pPr/w:spacing/@w:before')
    sa = p._p.xpath('./w:pPr/w:spacing/@w:after')
    line = p._p.xpath('./w:pPr/w:spacing/@w:line')
    
    # print summary
    if has_page_break or has_sectPr or txt:
        flags = []
        if has_page_break: flags.append("PAGE_BRK")
        if has_sectPr: flags.append("SECT_PR")
        flag_str = f"[{','.join(flags)}]" if flags else ""
        print(f"P{i:3d} (Pg~{page_num}) {flag_str:16s} jc={jc} sb={sb} sa={sa} | {txt[:70]}")
    
    if has_page_break or has_sectPr:
        page_num += 1
