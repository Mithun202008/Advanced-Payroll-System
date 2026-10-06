import docx
import sys
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')

path = r"C:\Users\HP\Downloads\PBL Report - Sample Template.docx"
doc = docx.Document(path)

print(f"=== SECTIONS ({len(doc.sections)}) ===")
for idx, s in enumerate(doc.sections):
    tm = round(s.top_margin.inches, 2) if s.top_margin else None
    bm = round(s.bottom_margin.inches, 2) if s.bottom_margin else None
    lm = round(s.left_margin.inches, 2) if s.left_margin else None
    rm = round(s.right_margin.inches, 2) if s.right_margin else None
    pw = round(s.page_width.inches, 2) if s.page_width else None
    ph = round(s.page_height.inches, 2) if s.page_height else None
    print(f"Sec {idx}: start={s.start_type}, top={tm}, btm={bm}, left={lm}, right={rm}, pw={pw}, ph={ph}")
    
    # Check headers and footers
    h_txts = [p.text.strip() for p in s.header.paragraphs if p.text.strip()]
    f_txts = [p.text.strip() for p in s.footer.paragraphs if p.text.strip()]
    f_xml = [p._p.xml for p in s.footer.paragraphs]
    has_page_field = any(('w:fldSimple' in x or 'PAGE' in x or 'w:instrText' in x) for x in f_xml)
    print(f"   Header: {h_txts}")
    print(f"   Footer: {f_txts} (has_page_num_field={has_page_field})")

print(f"\n=== ALL PARAGRAPHS ({len(doc.paragraphs)}) ===")
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    xml = p._p.xml
    has_br = 'w:br' in xml
    has_sect = bool(p._p.xpath('./w:pPr/w:sectPr'))
    jc = p._p.xpath('./w:pPr/w:jc/@w:val')
    sb = p._p.xpath('./w:pPr/w:spacing/@w:before')
    sa = p._p.xpath('./w:pPr/w:spacing/@w:after')
    line = p._p.xpath('./w:pPr/w:spacing/@w:line')
    
    # Check font of first run
    r0 = p.runs[0] if p.runs else None
    fn = r0.font.name if r0 and r0.font.name else 'None'
    sz = r0.font.size.pt if r0 and r0.font.size else 'None'
    bld = r0.bold if r0 else 'None'
    
    # Print headings or breaks or sample content
    is_interesting = (
        has_sect or has_br or 
        any(txt.startswith(x) for x in ['CHAPTER', '1.', '2.', '3.', '4.', '5.', '6.', '7.', '8.', 'ABSTRACT', 'TABLE OF CONTENTS', 'LIST OF', 'REFERENCES', 'APPENDIX', 'BONAFIDE', 'DECLARATION', 'ACKNOWLEDGEMENT', 'PROJECT REPORT', 'VISION', 'MISSION', 'PROGRAM', 'CIT', 'CHENNAI'])
    )
    if is_interesting or i < 40 or i > len(doc.paragraphs) - 30:
        flags = []
        if has_br: flags.append("BR")
        if has_sect: flags.append("SECT")
        flag_str = f"[{','.join(flags)}]" if flags else ""
        print(f"P{i:3d} {flag_str:10s} align={str(p.alignment):12s} jc={str(jc):10s} fn={fn} sz={sz} bld={bld} | {txt[:65]}")
