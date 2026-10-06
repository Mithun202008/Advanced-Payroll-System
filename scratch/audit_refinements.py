import docx
import pypdf
import sys

sys.stdout.reconfigure(encoding='utf-8')

doc = docx.Document('Enterprise_Payroll_Accounting_Management_System_PBL_REPORT.docx')

print("=== 1. AUDITING HEADING FONT SIZES ===")
chapter_sizes = []
section_sizes = []

for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if txt.startswith("CHAPTER ") or txt in ["REFERENCES", "APPENDIX"]:
        sz = [r.font.size.pt for r in p.runs if r.font.size]
        bold = [r.bold for r in p.runs]
        chapter_sizes.append((txt[:35], sz, bold, p.alignment))
    elif any(txt.startswith(f"{ch}.") for ch in range(1, 9)) and not txt[3:4].isdigit() and len(txt) < 80:
        sz = [r.font.size.pt for r in p.runs if r.font.size]
        bold = [r.bold for r in p.runs]
        section_sizes.append((txt[:35], sz, bold, p.alignment))

print(f"Found {len(chapter_sizes)} Chapter headings. Sample:")
for ch in chapter_sizes[:4]:
    print("  ", ch)

print(f"\nFound {len(section_sizes)} Section headings. Sample:")
for sec in section_sizes[:5]:
    print("  ", sec)

print("\n=== 2. AUDITING TABLE LEFT COLUMN BOLDING ===")
for t_idx, t in enumerate(doc.tables):
    col0_bolds = []
    col1_bolds = []
    # check first 3 data rows
    for r in t.rows[1:min(4, len(t.rows))]:
        c0_runs = r.cells[0].paragraphs[0].runs
        c1_runs = r.cells[1].paragraphs[0].runs if len(r.cells) > 1 else []
        c0_b = c0_runs[0].bold if c0_runs else False
        c1_b = c1_runs[0].bold if c1_runs else False
        col0_bolds.append(c0_b)
        col1_bolds.append(c1_b)
    title = t.rows[0].cells[0].text.strip()[:20]
    print(f"Table {t_idx:2d} ({len(t.rows):2d} rows): Header='{title}' | Col0 bold={col0_bolds} | Col1 bold={col1_bolds}")

print("\n=== 3. AUDITING PDF PAGE BREAKS & ALIGNMENT ===")
reader = pypdf.PdfReader('Enterprise_Payroll_Accounting_Management_System_PBL_REPORT.pdf')
print(f"Total PDF pages: {len(reader.pages)}")

checks = [
    (1, "Cover Page", lambda t: "Enterprise Payroll" in t and "October - 2026" in t),
    (2, "Vision/Mission Institute", lambda t: "Mission of the Institute:" in t),
    (3, "Vision/Mission Department", lambda t: "Vision of the Department:" in t),
    (4, "PEO & PO", lambda t: "PROGRAM EDUCATIONAL OBJECTIVES" in t),
    (5, "PSO", lambda t: "PROGRAM SPECIFIC OUTCOMES" in t),
    (6, "Bonafide Certificate", lambda t: "BONAFIDE CERTIFICATE" in t and "INTERNAL EXAMINER" in t),
    (7, "Declaration", lambda t: "DECLARATION" in t and "Date:" in t and "MITHUN A" in t),
    (8, "Acknowledgement", lambda t: "ACKNOWLEDGEMENT" in t and "[TEAMMATE NAME]" in t),
    (9, "Abstract", lambda t: "ABSTRACT" in t and "Keywords" in t),
    (10, "Table of Contents", lambda t: "TABLE OF CONTENTS" in t and "A.3 Self and Peer Assessment" in t),
    (11, "List of Figures", lambda t: "LIST OF FIGURES" in t),
    (12, "List of Tables", lambda t: "LIST OF TABLES" in t),
    (13, "List of Abbreviations", lambda t: "LIST OF ABBREVIATIONS" in t and "WA Washing Allowance" in t),
    (14, "Chapter 1 Introduction", lambda t: "CHAPTER 1: INTRODUCTION" in t and "1.1 Background" in t),
]

for page_no, name, cond in checks:
    txt = reader.pages[page_no - 1].extract_text() or ''
    res = cond(txt)
    status = "PASS" if res else "FAIL"
    print(f"Page {page_no:2d}: {name:30s} -> [{status}]")
