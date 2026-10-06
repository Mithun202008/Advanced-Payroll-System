import docx

docx_path = r"C:\Users\HP\OneDrive\Documents\Advanced_Payroll_System\PayRoll\Enterprise_Payroll_Accounting_Management_System_PBL_REPORT (2).docx"
doc = docx.Document(docx_path)

for i in [198, 249, 255, 260, 400, 405, 409, 414, 419, 422, 427, 432, 437, 441]:
    print(f"--- Around Drawing at [{i}] ---")
    for j in range(max(0, i-1), min(len(doc.paragraphs), i+4)):
        p = doc.paragraphs[j]
        has_draw = 'w:drawing' in p._p.xml
        t = repr(p.text)
        print(f"  [{j:03d}] draw={has_draw} text={t}")
