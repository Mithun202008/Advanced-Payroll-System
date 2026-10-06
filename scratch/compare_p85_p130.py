import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')
our_doc = docx.Document('Enterprise_Payroll_Accounting_Management_System_PBL_REPORT.docx')
ref_doc = docx.Document('pbl/RiskLens ( SYAMRIT , S HARSHAVARTHANAN ).docx')

print("=== OUR DOC P85 TO P130 ===")
for i in range(85, min(130, len(our_doc.paragraphs))):
    p = our_doc.paragraphs[i]
    sect = bool(p._p.xpath('./w:pPr/w:sectPr'))
    txt = p.text.strip()
    print(f"OUR P{i:3d}: sect={sect} runs={len(p.runs)} text={repr(txt[:60])}")

print("\n=== REF DOC P85 TO P130 ===")
for i in range(85, min(130, len(ref_doc.paragraphs))):
    p = ref_doc.paragraphs[i]
    sect = bool(p._p.xpath('./w:pPr/w:sectPr'))
    txt = p.text.strip()
    print(f"REF P{i:3d}: sect={sect} runs={len(p.runs)} text={repr(txt[:60])}")
