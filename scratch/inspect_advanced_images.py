import docx

doc = docx.Document('Advanced_Payroll_System_PBL_REPORT.docx')
print("Total paragraphs in Advanced docx:", len(doc.paragraphs))
for i, p in enumerate(doc.paragraphs):
    has_img = len(p._element.xpath('.//a:blip')) > 0 or len(p._element.xpath('.//w:drawing')) > 0
    if has_img:
        prev_t = doc.paragraphs[i-1].text[:60] if i > 0 else ""
        next_t = doc.paragraphs[i+1].text[:60] if i+1 < len(doc.paragraphs) else ""
        print(f"P{i} [IMAGE]: prev='{prev_t}' | next='{next_t}'")
    elif any(k in p.text for k in ['Figure ', 'Fig ', 'FIGURE ']):
        print(f"P{i} [CAPTION]: {p.text[:90]}")
