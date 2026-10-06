import docx
import os
import zipfile

docx_path = r"C:\Users\HP\OneDrive\Documents\Advanced_Payroll_System\PayRoll\Enterprise_Payroll_Accounting_Management_System_PBL_REPORT (2).docx"

doc = docx.Document(docx_path)

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

# Check images in zip
with zipfile.ZipFile(docx_path, 'r') as z:
    media_files = [f for f in z.namelist() if f.startswith('word/media/')]
    print(f"Total media files in docx: {len(media_files)}")
    for m in media_files:
        info = z.getinfo(m)
        print(f"  Media: {m} (size: {info.file_size} bytes)")

# Search for paragraphs with figures, screenshots, images, drawings
print("\n--- Paragraphs mentioning Figure, Screenshot, UI, or containing inline shapes ---")
for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    has_drawing = 'w:drawing' in p._p.xml
    if "Figure" in text or "Screenshot" in text or "Output" in text or "UI" in text or has_drawing:
        print(f"[{i}] (drawing={has_drawing}): {text[:100]}")
