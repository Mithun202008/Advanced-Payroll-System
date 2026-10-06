import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

docx_path = r"C:\Users\HP\OneDrive\Documents\Advanced_Payroll_System\PayRoll\Enterprise_Payroll_Accounting_Management_System_PBL_REPORT (2).docx"
doc = docx.Document(docx_path)

import xml.etree.ElementTree as ET

for i, p in enumerate(doc.paragraphs):
    if 'w:drawing' in p._p.xml:
        root = ET.fromstring(p._p.xml)
        extents = root.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent')
        cx, cy = 0, 0
        if extents:
            cx = int(extents[0].attrib.get('cx', 0))
            cy = int(extents[0].attrib.get('cy', 0))
        cx_in = cx / 914400.0
        cy_in = cy / 914400.0
        
        # Caption is usually next paragraph or 2 paragraphs down
        caption = ""
        for j in range(i+1, min(i+4, len(doc.paragraphs))):
            t = doc.paragraphs[j].text.strip()
            if t.startswith("Figure") or "Workflow" in t or "Architecture" in t or "Schema" in t:
                caption = t
                break
        
        print(f"Para [{i:03d}]: Width={cx_in:.2f}in, Height={cy_in:.2f}in | Align={p.alignment} | Indent={p.paragraph_format.left_indent} | Caption: {caption}")
