import docx

doc = docx.Document('Enterprise_Payroll_Accounting_Management_System_PBL_REPORT.docx')

print(f"Total tables: {len(doc.tables)}\n")
for t_idx, table in enumerate(doc.tables):
    rows = table.rows
    cols = table.columns
    print(f"=== TABLE {t_idx} ({len(rows)} rows, {len(cols)} cols) ===")
    
    # Print headers
    hdr_texts = [cell.text.strip().replace('\n', ' ') for cell in rows[0].cells]
    print(f"Headers: {hdr_texts}")
    
    # Print first 2-3 data rows
    for r_idx in range(1, min(4, len(rows))):
        row_texts = [cell.text.strip().replace('\n', ' ') for cell in rows[r_idx].cells]
        print(f"  Row {r_idx}: {row_texts}")
    print()
