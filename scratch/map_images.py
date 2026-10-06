import docx
import zipfile
import xml.etree.ElementTree as ET

docx_path = r"C:\Users\HP\OneDrive\Documents\Advanced_Payroll_System\PayRoll\Enterprise_Payroll_Accounting_Management_System_PBL_REPORT (2).docx"
doc = docx.Document(docx_path)

# Extract images to a temporary directory to inspect what they currently look like!
extract_dir = r"C:\Users\HP\OneDrive\Documents\Advanced_Payroll_System\PayRoll\scratch\extracted_doc_images"
import os
os.makedirs(extract_dir, exist_ok=True)

with zipfile.ZipFile(docx_path, 'r') as z:
    for f in z.namelist():
        if f.startswith('word/media/'):
            z.extract(f, r"C:\Users\HP\OneDrive\Documents\Advanced_Payroll_System\PayRoll\scratch\extracted_doc_images")

print("Extracted media to scratch/extracted_doc_images")

# Map paragraphs with drawings to image rIds and file names
rels = doc.part.rels
for i, p in enumerate(doc.paragraphs):
    if 'w:drawing' in p._p.xml:
        root = ET.fromstring(p._p.xml)
        blips = root.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
        rids = [b.attrib.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed') for b in blips]
        targets = [rels[r].target_ref for r in rids if r in rels]
        caption = doc.paragraphs[i+1].text if i+1 < len(doc.paragraphs) else ""
        if not caption.strip() and i+2 < len(doc.paragraphs):
            caption = doc.paragraphs[i+2].text
        print(f"Para [{i}]: Image target={targets} | Next Caption={caption[:80]}")
