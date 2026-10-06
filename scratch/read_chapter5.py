import docx

docx_path = r"C:\Users\HP\OneDrive\Documents\Advanced_Payroll_System\PayRoll\Enterprise_Payroll_Accounting_Management_System_PBL_REPORT (2).docx"
doc = docx.Document(docx_path)

for i in range(398, min(448, len(doc.paragraphs))):
    p = doc.paragraphs[i]
    has_drawing = 'w:drawing' in p._p.xml
    align = p.alignment
    indent = p.paragraph_format.left_indent
    print(f"[{i:03d}] (draw={has_drawing}, align={align}, ind={indent}): {p.text.strip()}")
