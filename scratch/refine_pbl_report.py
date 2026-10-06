import os
import sys
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
import win32com.client as win32
import pypdf

# Paths
TEMPLATE_PATH = r"C:\Users\HP\Downloads\PBL Report - Sample Template.docx"
OUTPUT_DOCX = os.path.abspath(r"Enterprise_Payroll_Accounting_Management_System_PBL_REPORT.docx")
OUTPUT_PDF = os.path.abspath(r"Enterprise_Payroll_Accounting_Management_System_PBL_REPORT.pdf")
DOC_IMAGES = os.path.abspath(r"scratch/doc_images")

print("1. Loading template:", TEMPLATE_PATH)
doc = Document(TEMPLATE_PATH)

bm_id_counter = 9000

def remove_existing_bookmark(doc, bm_name):
    body = doc._body._element
    b_starts = [el for el in body.iter() if el.tag.endswith('bookmarkStart') and el.attrib.get(qn('w:name')) == bm_name]
    for bs in b_starts:
        b_id = bs.attrib.get(qn('w:id'))
        p = bs.getparent()
        if p is not None:
            p.remove(bs)
        b_ends = [el for el in body.iter() if el.tag.endswith('bookmarkEnd') and el.attrib.get(qn('w:id')) == b_id]
        for be in b_ends:
            pe = be.getparent()
            if pe is not None:
                pe.remove(be)

def add_bookmark(p, bm_name):
    global bm_id_counter
    bm_id_counter += 1
    bm_id = str(bm_id_counter)
    remove_existing_bookmark(doc, bm_name)
    bm_start = parse_xml(f'<w:bookmarkStart {nsdecls("w")} w:id="{bm_id}" w:name="{bm_name}"/>')
    bm_end = parse_xml(f'<w:bookmarkEnd {nsdecls("w")} w:id="{bm_id}"/>')
    p._p.insert(0, bm_start)
    p._p.append(bm_end)

def add_body_p(doc, text, sa=4, ls=1.3, pbb=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = ls
    p.paragraph_format.space_after = Pt(sa)
    p.paragraph_format.space_before = Pt(0)
    if pbb:
        p.paragraph_format.page_break_before = True
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(13)
    run.font.bold = False
    return p

def add_bullet_p(doc, text, sa=3, ls=1.25):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = ls
    p.paragraph_format.space_after = Pt(sa)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.left_indent = Inches(0.25)
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(12.5)
    run.font.bold = False
    return p

def add_chapter_heading(doc, ch_num_str, ch_title_str, bm_name=None):
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p1.paragraph_format.page_break_before = True
    p1.paragraph_format.space_before = Pt(0)
    p1.paragraph_format.space_after = Pt(0)
    r1 = p1.add_run(ch_num_str)
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(16)
    r1.font.bold = True
    
    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(8)
    r2 = p2.add_run(ch_title_str)
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(16)
    r2.font.bold = True

    if bm_name:
        add_bookmark(p1, bm_name)
    return p1, p2

def add_section_heading(doc, text, bm_name=None, sb=6, sa=2, pbb=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_before = Pt(sb)
    p.paragraph_format.space_after = Pt(sa)
    if pbb:
        p.paragraph_format.page_break_before = True
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.bold = True
    if bm_name:
        add_bookmark(p, bm_name)
    return p

def add_subsection_heading(doc, text, bm_name=None, sb=5, sa=2, pbb=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_before = Pt(sb)
    p.paragraph_format.space_after = Pt(sa)
    if pbb:
        p.paragraph_format.page_break_before = True
    run = p.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(14)
    run.font.bold = True
    if bm_name:
        add_bookmark(p, bm_name)
    return p

def set_table_styling(table, col_widths):
    tblPr = table._tbl.tblPr
    # Remove existing borders
    for b in tblPr.xpath('./w:tblBorders'):
        tblPr.remove(b)
    borders_xml = f'''
    <w:tblBorders {nsdecls('w')}>
      <w:top w:val="single" w:sz="6" w:space="0" w:color="000000"/>
      <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>
      <w:bottom w:val="single" w:sz="6" w:space="0" w:color="000000"/>
      <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>
      <w:insideH w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
      <w:insideV w:val="single" w:sz="4" w:space="0" w:color="CCCCCC"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))
    
    # Set default cell margins: top=120, bottom=120, left=160, right=160 (dxa)
    for m in tblPr.xpath('./w:tblCellMar'):
        tblPr.remove(m)
    margins_xml = f'''
    <w:tblCellMar {nsdecls('w')}>
      <w:top w:w="120" w:type="dxa"/>
      <w:bottom w:w="120" w:type="dxa"/>
      <w:left w:w="160" w:type="dxa"/>
      <w:right w:w="160" w:type="dxa"/>
    </w:tblCellMar>
    '''
    tblPr.append(parse_xml(margins_xml))
    
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Row properties: tblHeader on row 0, cantSplit on all
    for r_idx, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if r_idx == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        
        # Explicit cell widths
        for c_idx, width in enumerate(col_widths):
            row.cells[c_idx].width = width
            tcPr = row.cells[c_idx]._tc.get_or_add_tcPr()
            # vertical centering
            for va in tcPr.xpath('./w:vAlign'):
                tcPr.remove(va)
            tcPr.append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>'))

def add_table_with_caption(doc, caption_str, col_widths, headers, data, col_alignments=None, bm_t_name=None, header_bg="EAEAEA", pbb=False):
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(8)
    p_cap.paragraph_format.space_after = Pt(3)
    if pbb:
        p_cap.paragraph_format.page_break_before = True
    run_cap = p_cap.add_run(caption_str)
    run_cap.font.name = "Times New Roman"
    run_cap.font.size = Pt(11.5)
    run_cap.font.bold = True
    if bm_t_name:
        add_bookmark(p_cap, bm_t_name)
        
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    set_table_styling(table, col_widths)
    
    # Headers - STRICTLY CENTERED, BOLD, WITH PADDING
    hdr_cells = table.rows[0].cells
    for idx, text in enumerate(headers):
        hdr_cells[idx].text = text
        p = hdr_cells[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.05
        for r in p.runs:
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)
            r.font.bold = True
        if header_bg:
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{header_bg}"/>')
            hdr_cells[idx]._tc.get_or_add_tcPr().append(shd)
            
    # Data rows
    for r_idx, row_data in enumerate(data):
        row = table.rows[r_idx + 1]
        for c_idx, val in enumerate(row_data):
            row.cells[c_idx].text = str(val)
            p = row.cells[c_idx].paragraphs[0]
            if col_alignments and c_idx < len(col_alignments):
                p.alignment = col_alignments[c_idx]
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.10
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(9.5)
                r.font.bold = False

    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(4)
    return table

def add_figure_with_caption(doc, img_name, caption_str, bm_f_name=None, width=Inches(4.2), pbb=False):
    img_path = os.path.join(DOC_IMAGES, img_name)
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(4)
    p_img.paragraph_format.space_after = Pt(1)
    if pbb:
        p_img.paragraph_format.page_break_before = True
    if os.path.exists(img_path):
        run_img = p_img.add_run()
        run_img.add_picture(img_path, width=width)
    else:
        print(f"WARNING: Image not found: {img_path}")
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.line_spacing = 1.15
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(5)
    run_cap = p_cap.add_run(caption_str)
    run_cap.font.name = "Times New Roman"
    run_cap.font.size = Pt(11.5)
    run_cap.font.bold = True
    if bm_f_name:
        add_bookmark(p_cap, bm_f_name)
    return p_cap

def add_code_listing(doc, title_str, code_str, pbb=False):
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_title.paragraph_format.space_before = Pt(4)
    p_title.paragraph_format.space_after = Pt(1)
    if pbb:
        p_title.paragraph_format.page_break_before = True
    r_t = p_title.add_run(title_str)
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(11)
    r_t.font.bold = True
    
    lines = code_str.strip().split('\n')
    for idx, line in enumerate(lines):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Pt(16)
        run = p.add_run(line)
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.bold = False
        
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(0)
    p_spacer.paragraph_format.space_after = Pt(3)

def make_pageref_cell(cell, bm_name, fallback_text="1", bold=False, font_size="22"):
    tcPr = cell._tc.get_or_add_tcPr()
    for va in tcPr.xpath('./w:vAlign'):
        tcPr.remove(va)
    tcPr.append(parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>'))
    
    p = cell.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    b_xml = "<w:b/>" if bold else ""
    p_xml = f'''
    <w:p xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
      <w:pPr>
        <w:spacing w:before="20" w:after="20"/>
        <w:jc w:val="center"/>
        <w:rPr>
          <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
          <w:sz w:val="{font_size}"/>
          {b_xml}
        </w:rPr>
      </w:pPr>
      <w:r>
        <w:rPr>
          <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
          <w:color w:val="000000"/>
          {b_xml}
        </w:rPr>
        <w:fldChar w:fldCharType="begin"/>
      </w:r>
      <w:r>
        <w:rPr>
          <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
          <w:color w:val="000000"/>
          {b_xml}
        </w:rPr>
        <w:instrText xml:space="preserve"> PAGEREF {bm_name} \\h </w:instrText>
      </w:r>
      <w:r>
        <w:rPr>
          <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
          <w:color w:val="000000"/>
          {b_xml}
        </w:rPr>
        <w:fldChar w:fldCharType="separate"/>
      </w:r>
      <w:r>
        <w:rPr>
          <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
          <w:color w:val="000000"/>
          <w:sz w:val="{font_size}"/>
          {b_xml}
        </w:rPr>
        <w:t>{fallback_text}</w:t>
      </w:r>
      <w:r>
        <w:rPr>
          <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
          <w:color w:val="000000"/>
          {b_xml}
        </w:rPr>
        <w:fldChar w:fldCharType="end"/>
      </w:r>
    </w:p>
    '''
    new_p = parse_xml(p_xml)
    cell._tc.replace(p._p, new_p)

print("Script base defined.")
