import docx

doc = docx.Document('Enterprise_Payroll_Accounting_Management_System_PBL_REPORT.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")

# Let's inspect paragraphs by chapter
current_ch = "FRONT MATTER"
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    if not txt:
        continue
    if txt.startswith("CHAPTER") or txt in ["REFERENCES", "APPENDIX", "ABSTRACT", "TABLE OF CONTENTS", "LIST OF TABLES", "LIST OF FIGURES", "LIST OF ABBREVIATIONS"]:
        current_ch = txt
        print(f"\n--- {current_ch} (P[{i}]) ---")
    elif len(txt) > 0 and (txt[0].isdigit() and len(txt) > 2 and txt[1] in ['.', ' ']):
        print(f"  Heading: {txt[:70]}")
    elif len(txt) > 0 and any(txt.startswith(x) for x in ["Table ", "Figure "]):
        print(f"  Caption: {txt[:70]}")
