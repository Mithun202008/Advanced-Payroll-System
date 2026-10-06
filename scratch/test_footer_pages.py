import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(1.0)
sec.bottom_margin = Inches(1.0)
sec.left_margin = Inches(1.0)
sec.right_margin = Inches(1.0)

# Configure footer
footer = sec.footer
footer.is_linked_to_previous = False
p_ftr = footer.paragraphs[0]
p_ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_ftr.text = ""

# Add PAGE field properly
def add_page_number_to_paragraph(p):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    
    r1 = p.add_run()
    r1.font.name = "Times New Roman"
    r1.font.size = Pt(11)
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    r1._r.append(fldChar1)
    
    r2 = p.add_run()
    r2.font.name = "Times New Roman"
    r2.font.size = Pt(11)
    instr = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    r2._r.append(instr)
    
    r3 = p.add_run()
    r3.font.name = "Times New Roman"
    r3.font.size = Pt(11)
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    r3._r.append(fldChar2)
    
    r4 = p.add_run("1") # fallback default value
    r4.font.name = "Times New Roman"
    r4.font.size = Pt(11)
    
    r5 = p.add_run()
    r5.font.name = "Times New Roman"
    r5.font.size = Pt(11)
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    r5._r.append(fldChar3)

add_page_number_to_paragraph(p_ftr)

p = doc.add_paragraph("This is Page 1 content")
doc.add_page_break()
doc.add_paragraph("This is Page 2 content")
doc.add_page_break()
doc.add_paragraph("This is Page 3 content")

doc.save("scratch/test_footer_pages.docx")
print("Saved scratch/test_footer_pages.docx")
