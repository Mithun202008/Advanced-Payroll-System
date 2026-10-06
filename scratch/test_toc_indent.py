import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

doc = Document()
sec = doc.sections[0]
sec.left_margin = Inches(1.0)
sec.right_margin = Inches(1.0)

tab_xml = f'<w:tabs {nsdecls("w")}><w:tab w:val="right" w:leader="dot" w:pos="9028"/></w:tabs>'

p1 = doc.add_paragraph()
p1.alignment = WD_ALIGN_PARAGRAPH.LEFT
p1._p.get_or_add_pPr().append(parse_xml(tab_xml))
r1 = p1.add_run("CHAPTER 1: INTRODUCTION\t14")
r1.bold = True

p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.LEFT
p2.paragraph_format.left_indent = Inches(0.25)
p2._p.get_or_add_pPr().append(parse_xml(tab_xml))
r2 = p2.add_run("1.1 Background\t14")

doc.save("scratch/test_toc_indent.docx")
print("Saved scratch/test_toc_indent.docx successfully!")
