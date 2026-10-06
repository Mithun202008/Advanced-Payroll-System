import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import win32com.client as win32
import os, pypdf

doc = Document()
s = doc.sections[0]
s.top_margin = Inches(1.0)
s.bottom_margin = Inches(1.0)
s.left_margin = Inches(1.0)
s.right_margin = Inches(1.0)

# Add TOC title centered
p_toc = doc.add_paragraph()
p_toc.alignment = WD_ALIGN_PARAGRAPH.CENTER
p_toc.paragraph_format.space_before = Pt(6)
p_toc.paragraph_format.space_after = Pt(6)
r = p_toc.add_run("TABLE OF CONTENTS")
r.bold = True
r.font.size = Pt(14)

tab_xml = f'<w:tabs {nsdecls("w")}><w:tab w:val="right" w:leader="dot" w:pos="9028"/></w:tabs>'

for i in range(44):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    p._p.get_or_add_pPr().append(parse_xml(tab_xml))
    r1 = p.add_run(f"Chapter or Section Heading {i+1}")
    r1.font.size = Pt(10.5)
    r2 = p.add_run(f"\t{i+14}")
    r2.font.size = Pt(10.5)

doc.save("scratch/test_toc_single_page.docx")

word = win32.Dispatch("Word.Application")
word.Visible = False
d = word.Documents.Open(os.path.abspath("scratch/test_toc_single_page.docx"))
d.SaveAs(os.path.abspath("scratch/test_toc_single_page.pdf"), FileFormat=17)
d.Close()
word.Quit()

reader = pypdf.PdfReader("scratch/test_toc_single_page.pdf")
print("TOC page count with 44 rows at 10.5pt:", len(reader.pages))
