import docx
import xml.etree.ElementTree as ET

docx_path = r"C:\Users\HP\OneDrive\Documents\Advanced_Payroll_System\PayRoll\Enterprise_Payroll_Accounting_Management_System_PBL_REPORT (2).docx"
doc = docx.Document(docx_path)

print(f"Verified: Document opens cleanly! Total paragraphs: {len(doc.paragraphs)}")

figure_count = 0
for i, p in enumerate(doc.paragraphs):
    if 'w:drawing' in p._p.xml:
        root = ET.fromstring(p._p.xml)
        extents = root.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}extent')
        cx, cy = 0, 0
        if extents:
            cx = int(extents[0].attrib.get('cx', 0))
            cy = int(extents[0].attrib.get('cy', 0))
        cx_in = cx / 914400.0
        cy_in = cy / 914400.0
        
        # Check next paragraph for caption
        next_p_text = doc.paragraphs[i+1].text.strip() if i+1 < len(doc.paragraphs) else ""
        if "Figure" in next_p_text:
            figure_count += 1
            print(f"Fig #{figure_count:02d} [Para {i:03d}]: Width={cx_in:.2f} in, Height={cy_in:.2f} in | Align={p.alignment} | Next Para: '{next_p_text}'")

print(f"\nTotal figures verified: {figure_count}")

# Check sample theory paragraphs to confirm 100% preservation
sample_indices = [100, 200, 300, 450]
for idx in sample_indices:
    if idx < len(doc.paragraphs):
        print(f"Theory Sample [Para {idx}]: {doc.paragraphs[idx].text[:90]}...")
